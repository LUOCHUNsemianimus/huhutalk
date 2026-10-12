#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""HuhuTalk 手机手册（ManualPage.java）⇄ Word 往返工具。

需求来源：手册内容过时了，需求方要在 Word 里改，改完再写回模组。

用法
----
    python tools/manual_docx.py export                # 生成 docs/manual-entries.docx
    python tools/manual_docx.py export -o X.docx      # 指定输出
    python tools/manual_docx.py import --dry-run      # 只看会改什么（打印改动清单 + diff，不写盘）
    python tools/manual_docx.py import                # 写回 ManualPage.java（先备份 .bak）
    python tools/manual_docx.py import --allow-delete # 允许删掉"文档里没有的"条目

往返约定（导出的 Word 里也写着同一份）
------------------------------------
1. **不要改方括号里的编号**（`[2]` = 大栏目、`[2-3]` = 第 2 节第 3 条）。标题和正文随便改。
2. 想新增一条：复制一条「标题行 + 正文」，编号写成 `[新]` 或不写编号；它会被追加到所在大栏目的末尾。
3. 想删掉一条：整段删掉，import 时加 `--allow-delete`；不加则保留并在改动清单里告警。
4. 正文里一个自然段 = 一行（对应模组里的 `\\n` 换行），空段落保留为空行。

实现上的两个刻意选择
------------------
* **只改该改的语句**：`sections.add(sec(...))` 里没有变化的语句整段原样搬回去（连它原来的
  字符串拼接换行都保留），所以 diff 只落在你真正编辑过的条目上，而不是把 42 条全重排版。
* **写盘前自校验**：生成新源码后立刻重新解析一遍，把"解析出来的内容"和"我想要的内容"逐条对比，
  不一致就**拒绝写盘**。这样转义/拼接出问题时不会静默改坏手册文案。
"""

from __future__ import annotations

import argparse
import datetime
import difflib
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MANUAL_JAVA = ROOT / "src" / "main" / "java" / "com" / "huhutalk" / "phone" / "pages" / "ManualPage.java"
DEFAULT_DOCX = ROOT / "docs" / "manual-entries.docx"

STMT_INDENT = " " * 8          # sections.add(sec( 这一层的缩进
ENTRY_INDENT = " " * 16        # entry( 这一层的缩进
CONT_INDENT = " " * 24         # 字符串拼接续行 "+" 的缩进


def info(msg: str = "") -> None:
    print(msg)


def die(msg: str) -> None:
    print("ERROR: " + msg, file=sys.stderr)
    sys.exit(2)


# ============================================================================
# Java 侧：屏蔽注释 / 扫描 / 解析
# ============================================================================


def mask_comments(src: str) -> str:
    """把注释替换成等长空白（所有偏移不变），字符串字面量原样保留。"""
    out = []
    i, n = 0, len(src)
    while i < n:
        c = src[i]
        nxt = src[i + 1] if i + 1 < n else ""
        if c == "/" and nxt == "/":
            j = src.find("\n", i)
            j = n if j < 0 else j
            out.append(" " * (j - i))
            i = j
        elif c == "/" and nxt == "*":
            j = src.find("*/", i + 2)
            j = n if j < 0 else j + 2
            out.append("".join("\n" if ch == "\n" else " " for ch in src[i:j]))
            i = j
        elif c == '"':
            j = i + 1
            while j < n:
                if src[j] == "\\":
                    j += 2
                    continue
                if src[j] == '"':
                    j += 1
                    break
                j += 1
            out.append(src[i:j])
            i = j
        else:
            out.append(c)
            i += 1
    return "".join(out)


_ESCAPES = {"n": "\n", "t": "\t", "r": "\r", '"': '"', "\\": "\\", "'": "'"}


class Scanner:
    """只认字符串字面量、`+`、逗号和括号的极简扫描器（坐标基于 mask_comments 后的文本）。"""

    def __init__(self, s: str, i: int = 0) -> None:
        self.s = s
        self.i = i

    def ws(self) -> None:
        while self.i < len(self.s) and self.s[self.i] in " \t\r\n":
            self.i += 1

    def take(self, lit: str) -> bool:
        self.ws()
        if self.s.startswith(lit, self.i):
            self.i += len(lit)
            return True
        return False

    def expect(self, lit: str) -> None:
        if not self.take(lit):
            raise ValueError(f"期望 {lit!r}，实际 {self.s[self.i:self.i + 60]!r}")

    def peek(self) -> str:
        self.ws()
        return self.s[self.i] if self.i < len(self.s) else ""

    def read_string_expr(self) -> tuple[str, int, int]:
        """读 `"a" + "b"`，返回 (解码后的内容, 起, 止)。"""
        start = self.i
        parts: list[str] = []
        while True:
            self.ws()
            if start == self.i:
                start = self.i
            if self.peek() != '"':
                raise ValueError(f"期望字符串字面量，实际 {self.s[self.i:self.i + 60]!r}")
            j = self.i + 1
            buf: list[str] = []
            while True:
                if j >= len(self.s):
                    raise ValueError("字符串没有收尾的引号")
                c = self.s[j]
                if c == "\\":
                    buf.append(_ESCAPES.get(self.s[j + 1], self.s[j + 1]))
                    j += 2
                    continue
                if c == '"':
                    break
                buf.append(c)
                j += 1
            parts.append("".join(buf))
            self.i = j + 1
            self.ws()
            if self.peek() == "+":
                self.i += 1
                continue
            break
        return "".join(parts), start, self.i

    def read_image_expr(self) -> tuple[str | None, str | None, int, int]:
        """读 `entry(...)` 的**可选第三个参数**（配图）。

        支持两种写法（都出现在真实源码里）：
          * 字符串字面量：`"manual_fig_1"`（老写法）
          * 源码辅助函数：`fig(1)`（手工写回手册时用的，定义在 ManualPage 里）

        @return (值, 原样文本, 起, 止)。值为 None = 没有配图（参数写成 null / 空串）。
        """
        self.ws()
        start = self.i
        if self.peek() == '"':
            val, s, e = self.read_string_expr()
            return (val if val else None), None, s, e
        depth = 0
        while self.i < len(self.s):
            c = self.s[self.i]
            if c == "(":
                depth += 1
            elif c == ")":
                if depth == 0:
                    break
                depth -= 1
            elif c == "," and depth == 0:
                break
            self.i += 1
        raw = self.s[start:self.i].strip()
        if raw in ("null", "NULL", ""):
            return None, raw, start, self.i
        m = re.fullmatch(r"fig\(\s*(\d+)\s*\)", raw)
        if m:
            return f"manual_fig_{m.group(1)}", raw, start, self.i
        return raw, raw, start, self.i


def parse_manual(src: str) -> dict:
    """解析 `sections.add(sec(...));` 区域，并记录每条语句/每个字面量的源码区间。"""
    masked = mask_comments(src)
    sections: list[dict] = []
    marker = "sections.add(sec("
    pos = 0
    first_stmt_line_start = -1
    last_end = -1
    prev_end = -1
    while True:
        k = masked.find(marker, pos)
        if k < 0:
            break
        line_start = masked.rfind("\n", 0, k) + 1
        if first_stmt_line_start < 0:
            first_stmt_line_start = line_start
        # 语句之间的原文（注释、空行、缩进）也要留住 —— 不然回写时我写在两条 sections.add 之间的
        # 说明性注释会被悄悄删掉。第一句的间隙是空串（span 从它所在行的行首开始）。
        gap = "" if prev_end < 0 else src[prev_end:line_start]
        sc = Scanner(masked, k + len(marker))
        title, title_start, title_end = sc.read_string_expr()
        if sc.peek() == ",":
            sc.i += 1
        entries: list[dict] = []
        insert_at = -1
        while True:
            sc.ws()
            if masked[sc.i] == ")":
                insert_at = sc.i
                break
            entry_start = sc.i
            if not sc.take("entry"):
                raise ValueError(f"sec 里出现了非 entry 的表达式：{masked[sc.i:sc.i + 60]!r}")
            sc.expect("(")
            e_title, et_s, et_e = sc.read_string_expr()
            sc.expect(",")
            e_body, eb_s, eb_e = sc.read_string_expr()
            e_image = None
            e_raw_image = None
            if sc.peek() == ",":
                sc.i += 1
                e_image, e_raw_image, _, _ = sc.read_image_expr()
            sc.expect(")")
            entries.append({
                "title": e_title, "body": e_body, "image": e_image,
                "title_span": (et_s, et_e), "body_span": (eb_s, eb_e),
                "entry_span": (entry_start, sc.i),
                "raw_title": src[et_s:et_e], "raw_body": src[eb_s:eb_e],
                "raw_image": e_raw_image,
            })
            if sc.peek() == ",":
                sc.i += 1
                continue
            break
        sc.expect(")")   # 收 sec(
        sc.expect(")")   # 收 sections.add(
        sc.take(";")
        last_end = sc.i
        pos = sc.i
        prev_end = last_end
        sections.append({
            "src_id": len(sections) + 1,
            "title": title,
            "title_span": (title_start, title_end),
            "raw_title": src[title_start:title_end],
            "stmt_span": (line_start, last_end),
            "gap": gap,
            "insert_at": insert_at,
            "entries": entries,
        })
    if first_stmt_line_start < 0:
        raise ValueError("在 ManualPage.java 里没找到 sections.add(sec( 块")
    return {"sections": sections, "span": (first_stmt_line_start, last_end)}


def java_escape(text: str) -> str:
    out = []
    for ch in text:
        if ch == "\\":
            out.append("\\\\")
        elif ch == '"':
            out.append('\\"')
        elif ch == "\n":
            out.append("\\n")
        elif ch == "\r":
            out.append("\\r")
        elif ch == "\t":
            out.append("\\t")
        else:
            out.append(ch)
    return "".join(out)


def java_body_expr(body: str) -> str:
    """正文 → Java 字符串表达式（一段一行，与文件现有写法一致）。"""
    lines = body.split("\n")
    if len(lines) == 1:
        return '"' + java_escape(lines[0]) + '"'
    parts = ['"' + java_escape(line) + '\\n"' for line in lines[:-1]]
    parts.append('"' + java_escape(lines[-1]) + '"')
    expr = parts[0]
    for extra in parts[1:]:
        expr += "\n" + CONT_INDENT + "+ " + extra
    return expr


def render_stmt(sec: dict) -> str:
    """重写一条 `sections.add(sec(...));`（只在真的改了东西时才调用）。

    没动过的条目沿用它**原来的字面量写法**（连中途换行都保持），
    这样 diff 只会落在真正编辑过的那一条上。
    """
    dirty = bool(
        sec.get("dirty")
        or sec.get("title_dirty")
        or any(e.get("dirty") or e.get("new") for e in sec["entries"])
    )
    if not dirty and sec.get("raw_stmt"):
        return sec["raw_stmt"]

    head = f'{STMT_INDENT}sections.add(sec({sec["raw_title"]},'
    if sec.get("title_dirty") or not sec.get("raw_title"):
        head = f'{STMT_INDENT}sections.add(sec({java_quote(sec["title"])},'

    lines = [head]
    entries = sec["entries"]
    if entries:
        for idx, en in enumerate(entries):
            tail = ")));" if idx == len(entries) - 1 else "),"
            title_expr = java_quote(en["title"]) if en.get("title_dirty") or en.get("new") else en["raw_title"]
            body_expr = java_body_expr(en["body"]) if en.get("body_dirty") or en.get("new") else en["raw_body"]
            if (en.get("new") or en.get("dirty") or en.get("title_dirty")
                    or en.get("body_dirty") or en.get("image_dirty")):
                img_expr = image_expr(en)
            else:
                img_expr = en.get("raw_image")   # 没动过的条目沿用原来的写法（fig(1) 或 "manual_fig_1"）
            args = f"{title_expr}, {body_expr}"
            if img_expr:
                args += f", {img_expr}"
            lines.append(f'{ENTRY_INDENT}entry({args}{tail}')
    else:
        lines = [f'{STMT_INDENT}sections.add(sec({java_quote(sec["title"])}));']
    return "\n".join(lines)


def java_quote(text: str) -> str:
    return '"' + java_escape(text) + '"'


def image_expr(en: dict) -> str | None:
    """配图值 → Java 表达式：`manual_fig_N` 写成 `fig(N)`，其它键写成字符串字面量，无图返回 None。"""
    key = en.get("image")
    if not key:
        return None
    m = re.fullmatch(r"manual_fig_(\d+)", key)
    return f"fig({m.group(1)})" if m else java_quote(key)


def figure_label(key: str | None) -> str:
    """配图键 → Word 里给人看的名字（`manual_fig_3` → `图3`）。"""
    if not key:
        return "无"
    m = re.fullmatch(r"manual_fig_(\d+)", key)
    return f"图{m.group(1)}" if m else key


def figure_key(label: str) -> str:
    """Word 里的名字 → 配图键（`图3` → `manual_fig_3`；其它原样当键）。"""
    m = re.fullmatch(r"图\s*(\d+)", label.strip())
    return f"manual_fig_{m.group(1)}" if m else label.strip()


def model_of(parsed: dict) -> list[tuple]:
    """仅内容（大栏目名 + 条目名/正文）的快照，用于写盘前自校验。"""
    return [
        (sec["title"], [(e["title"], e["body"]) for e in sec["entries"]])
        for sec in parsed["sections"]
    ]


# ============================================================================
# Word 侧
# ============================================================================

CJK_FONT = "微软雅黑"

PREAMBLE = [
    "【改法说明 · 留着不影响我回写】",
    "1. 方括号里的编号（[2] 是大栏目、[2-3] 是第 2 节第 3 条）是回写时的定位标记，请不要改、不要删；"
    "标题文字和正文随你改。正文一个自然段 = 一行（对应游戏里的一次换行）。",
    "2. 新增一条：复制任意一条的「标题行 + 正文」，编号写成 [新]（或不写编号），它会被追加到所在大栏目的末尾。",
    "3. 删掉一条：整段删掉；我回写时会先列清单告警，确认后才真删。",
    "4. 大栏目顺序 = 手册首页顺序；条目顺序 = 点进去后的列表顺序，想调就整块上下移动。",
    "5. 别用 Word 的自动更正去改引号（中文引号请保持“”），也别在正文里用真实的换行符排版 —— 换行＝新段落＝模组里的一行。",
    "6. 条目末尾那一行【配图：图N】= 这条正文下方挂的图（图放在 assets/huhutalk/manual/，随模组打包）。"
    "想换图就改编号、想去掉就整行删掉或写成【配图：无】、想加就照抄一行。",
]


def _set_cjk(style) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = CJK_FONT
    try:
        style.font.size = Pt(10.5)
    except Exception:
        pass
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), CJK_FONT)
    rfonts.set(qn("w:hAnsi"), CJK_FONT)


def external_sections() -> list[tuple[str, str]]:
    """扫工作区里的 `registerSection(`（附属模组往手册追加的栏目），只用于文末附录。"""
    found: list[tuple[str, str]] = []
    for path in ROOT.rglob("*.java"):
        if path.name == "ManualPage.java":
            continue
        try:
            text = path.read_text(encoding="utf-8")
        except OSError:
            continue
        for m in re.finditer(r'registerSection\(\s*"([^"]+)"\s*,\s*"([^"]+)"', text):
            found.append((m.group(1), m.group(2)))
    return found


def do_export(docx_path: Path) -> int:
    from docx import Document

    parsed = parse_manual(MANUAL_JAVA.read_text(encoding="utf-8"))
    sections = parsed["sections"]

    doc = Document()
    _set_cjk(doc.styles["Normal"])
    for name in ("Title", "Heading 1", "Heading 2"):
        if name in [s.name for s in doc.styles]:
            _set_cjk(doc.styles[name])

    doc.add_heading("HuhuTalk 手机手册 · 条目文本（可编辑版）", level=0)
    doc.add_paragraph("来源：src/main/java/com/huhutalk/phone/pages/ManualPage.java 的内置栏目"
                      "（附属模组追加的栏目不在本文件驱动范围内，见文末附录）。")
    for line in PREAMBLE:
        doc.add_paragraph(line)

    for sec in sections:
        doc.add_heading(f'[{sec["src_id"]}] 大栏目：{sec["title"]}', level=1)
        for ei, en in enumerate(sec["entries"], 1):
            doc.add_heading(f'[{sec["src_id"]}-{ei}] {en["title"]}', level=2)
            for line in en["body"].split("\n"):
                doc.add_paragraph(line)
            if en.get("image"):
                doc.add_paragraph(f'【配图：{figure_label(en["image"])}】')

    doc.add_page_break()
    doc.add_heading("附：由附属模组追加的栏目（不在本文件驱动范围内）", level=1)
    ext = external_sections()
    if ext:
        doc.add_paragraph("下面这些栏目是别的模组在运行时调 ManualPage.registerSection(...) 追加的，"
                          "文案在各自的模组里，改这份文档没用：")
        for mod_id, title in ext:
            doc.add_paragraph(f"· {title}（来自 {mod_id}）")
    else:
        doc.add_paragraph("（当前没有扫到 registerSection 调用。）")

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    info(f"已导出 {docx_path}")
    info(f"  大栏目 {len(sections)} 节、条目 {sum(len(s['entries']) for s in sections)} 条")
    return 0


HEAD_SEC_RE = re.compile(r"^\[(\d+)\]")
HEAD_ENTRY_RE = re.compile(r"^\[(\d+)-(\d+)\]")
IMG_DIRECTIVE_RE = re.compile(r"^【配图[：:]\s*(.*?)】$")


def _heading_level(para) -> int:
    """段落是不是 Heading N（1/2）。

    ⚠ 不要写 `style.name == "Heading 1"` 这种死等式：Word 另存之后样式名可能变成
    `heading 1`（小写）/ styleId `Heading1`，两种都得认，否则玩家一存就"读不到条目"。
    """
    st = getattr(para, "style", None)
    if st is None:
        return 0
    name = (st.name or "").strip().lower().replace(" ", "")
    sid = (getattr(st, "style_id", "") or "").strip().lower()
    for lvl in (1, 2, 3):
        if name == f"heading{lvl}" or sid == f"heading{lvl}":
            return lvl
    return 0


def read_docx(docx_path: Path) -> list[dict]:
    """把 Word 读回成 [{src_id, title, entries:[{src_id, title, body}]}]（只认带编号的标题）。"""
    from docx import Document

    doc = Document(str(docx_path))
    sections: list[dict] = []
    cur_sec = None
    cur_entry = None
    for para in doc.paragraphs:
        text = para.text.strip()
        lvl = _heading_level(para)
        if lvl == 1:
            m = HEAD_SEC_RE.match(text)
            if not m:
                cur_sec = None       # 附录之类：从这里开始不再收集
                cur_entry = None
                continue
            title = text[m.end():].strip().lstrip("：: ").strip()
            if title.startswith("大栏目"):
                title = title[len("大栏目"):].strip().lstrip("：: ").strip()
            cur_sec = {"src_id": int(m.group(1)), "title": title, "entries": []}
            sections.append(cur_sec)
            cur_entry = None
            continue
        if lvl == 2:
            if cur_sec is None:
                continue
            m = HEAD_ENTRY_RE.match(text)
            # 没有合法编号（例如用户写的 "[新] xxx" 或干脆不写）时，把开头任意方括号标记剥掉再当标题
            title = text[m.end():].strip() if m else re.sub(r"^\[[^\]]*\]\s*", "", text).strip()
            cur_entry = {
                "src_id": (int(m.group(1)), int(m.group(2))) if m else None,
                "title": title,
                "lines": [],
                "image": None,
                "has_directive": False,
            }
            cur_sec["entries"].append(cur_entry)
            continue
        if cur_entry is not None:
            mimg = IMG_DIRECTIVE_RE.match(text)
            if mimg:
                raw = mimg.group(1).strip()
                cur_entry["image"] = None if raw in ("无", "-", "none", "None", "") else figure_key(raw)
                cur_entry["has_directive"] = True
                continue
            cur_entry["lines"].append(para.text)
    for sec in sections:
        for en in sec["entries"]:
            lines = en.pop("lines")
            while lines and lines[-1].strip() == "":
                lines.pop()
            while lines and lines[0].strip() == "":
                lines.pop(0)
            en["body"] = "\n".join(lines)
    return sections


def do_import(docx_path: Path, dry_run: bool, allow_delete: bool) -> int:
    src = MANUAL_JAVA.read_text(encoding="utf-8")
    parsed = parse_manual(src)
    old_sections = parsed["sections"]
    doc_sections = read_docx(docx_path)
    if not doc_sections:
        die(f"在 {docx_path} 里没读到任何带编号的标题 —— 编号（[2]、[2-3]）被改掉了吗？")

    by_id = {s["src_id"]: s for s in old_sections}
    changes: list[str] = []
    merged: list[dict] = []

    for doc_sec in doc_sections:
        sid = doc_sec["src_id"]
        base = by_id.get(sid)
        if base is None:
            new_sec = {"src_id": sid, "title": doc_sec["title"], "raw_title": java_quote(doc_sec["title"]),
                       "entries": [], "raw_stmt": "", "title_dirty": True}
            for en in doc_sec["entries"]:
                new_sec["entries"].append({"title": en["title"], "body": en["body"], "image": en.get("image"),
                                           "new": True, "raw_title": "", "raw_body": "", "raw_image": None,
                                           "title_dirty": True, "body_dirty": True})
            merged.append(new_sec)
            changes.append(f"新增大栏目 [{sid}] {doc_sec['title']}（{len(doc_sec['entries'])} 条）")
            continue

        sec = dict(base)
        sec["title_dirty"] = base["title"] != doc_sec["title"]
        sec["title"] = doc_sec["title"]
        if sec["title_dirty"]:
            changes.append(f"大栏目 [{sid}] 改名：{base['title']} → {doc_sec['title']}")
        sec["raw_stmt"] = src[base["stmt_span"][0]:base["stmt_span"][1]]

        pool = [dict(e) | {"dirty": False, "new": False} for e in base["entries"]]
        kept: list[dict] = []
        seen: set[int] = set()
        extras: list[dict] = []
        for en in doc_sec["entries"]:
            eid = en["src_id"]
            if eid is not None and 1 <= eid[1] <= len(pool):
                orig = pool[eid[1] - 1]
                seen.add(eid[1])
                item = dict(orig)
                item["title_dirty"] = orig["title"] != en["title"]
                item["body_dirty"] = orig["body"] != en["body"]
                if item["title_dirty"]:
                    changes.append(f'  [{sid}-{eid[1]}] 标题：{orig["title"]} → {en["title"]}')
                if item["body_dirty"]:
                    changes.append(f'  [{sid}-{eid[1]}] 正文已改（{len(orig["body"])} → {len(en["body"])} 字）：{en["title"]}')
                item["title"] = en["title"]
                item["body"] = en["body"]
                if en.get("has_directive"):
                    item["image_dirty"] = (orig.get("image") or None) != (en.get("image") or None)
                    item["image"] = en["image"]
                    if item["image_dirty"]:
                        changes.append(f'  [{sid}-{eid[1]}] 配图：{figure_label(orig.get("image"))} → {figure_label(en.get("image"))}')
                item["dirty"] = item["title_dirty"] or item["body_dirty"] or item.get("image_dirty", False)
                kept.append(item)
                continue

            # 没有合法编号（文档里写成 [新] 之类）：先按"标题完全相同"认领一条还没被认领的旧条目，
            # 这样**同一份文档回写两次是幂等的**（第一次追加的条目第二次会被认成同一条而不是再追加一条）。
            match = next((k for k, o in enumerate(pool) if (k + 1) not in seen and o["title"] == en["title"]), None)
            if match is not None:
                orig = pool[match]
                seen.add(match + 1)
                item = dict(orig)
                item["title_dirty"] = False
                item["body_dirty"] = orig["body"] != en["body"]
                if item["body_dirty"]:
                    changes.append(f'  [{sid}-{match + 1}]（无编号，按标题认领）正文已改：{en["title"]}')
                item["body"] = en["body"]
                if en.get("has_directive"):
                    item["image_dirty"] = (orig.get("image") or None) != (en.get("image") or None)
                    item["image"] = en["image"]
                item["dirty"] = item["body_dirty"] or item.get("image_dirty", False)
                kept.append(item)
                continue

            extras.append({"title": en["title"], "body": en["body"], "image": en.get("image"),
                           "new": True, "title_dirty": True, "body_dirty": True,
                           "raw_title": "", "raw_body": "", "raw_image": None})
            changes.append(f'  在 [{sid}] 末尾新增条目：{en["title"]}')
        for ei, orig in enumerate(base["entries"], 1):
            if ei in seen:
                continue
            if allow_delete:
                changes.append(f'  [{sid}-{ei}] 删除条目：{orig["title"]}')
            else:
                kept.append(dict(orig) | {"dirty": False, "new": False})
                changes.append(f'  ⚠ [{sid}-{ei}] {orig["title"]} 在文档里找不到，已保留（要真删请加 --allow-delete）')
        kept.extend(extras)
        sec["entries"] = kept
        # ⚠ 必须显式标脏"增删过条目"：否则 render_stmt 会把原文整段搬回去，
        #   被删掉的条目会复活（自校验抓到过一次，别去掉这一行）。
        sec["dirty"] = bool(extras) or len(kept) != len(base["entries"])
        merged.append(sec)

    doc_ids = {s["src_id"] for s in doc_sections}
    for base in old_sections:
        if base["src_id"] not in doc_ids:
            merged.append(dict(base) | {"title_dirty": False,
                                        "entries": [dict(e) | {"dirty": False, "new": False} for e in base["entries"]]})
            changes.append(f'⚠ 大栏目 [{base["src_id"]}] {base["title"]} 在文档里找不到，已保留')

    # ---- 生成 ----
    # 每句前面带上它**原来**的间隙原文（注释/空行），于是"没动过的语句连注释一起原样保留"。
    new_stmts = [sec.get("gap", "\n") + render_stmt(sec) for sec in merged]
    new_block = "".join(new_stmts)
    new_src = src[:parsed["span"][0]] + new_block + src[parsed["span"][1]:]

    # ---- 写盘前自校验：重新解析，内容必须与预期逐条一致 ----
    try:
        reparsed = parse_manual(new_src)
    except Exception as exc:  # noqa: BLE001
        die(f"生成的源码解析失败（拒绝写盘）：{exc}")
    want = [(s["title"], [(e["title"], e["body"]) for e in s["entries"]]) for s in merged]
    got = model_of(reparsed)
    if want != got:
        die("生成的源码重新解析后与预期不一致（拒绝写盘）：\n"
            f"  预期 {len(want)} 节 / {sum(len(x[1]) for x in want)} 条\n"
            f"  实际 {len(got)} 节 / {sum(len(x[1]) for x in got)} 条")

    if not changes:
        info("没有检测到任何改动。")
    else:
        info(f"检测到 {len(changes)} 处改动：")
        for c in changes:
            info("  " + c)

    if dry_run:
        diff = list(difflib.unified_diff(src.splitlines(), new_src.splitlines(),
                                        fromfile="ManualPage.java(now)", tofile="ManualPage.java(after)",
                                        lineterm="", n=2))
        if diff:
            info()
            info("--- diff ---")
            for line in diff[:200]:
                print(line)
            if len(diff) > 200:
                info(f"... 还有 {len(diff) - 200} 行")
        else:
            info("dry-run：源码不会变（幂等 ✓）")
        return 0

    if new_src == src:
        info("源码无需改写（内容与格式都一致）。")
        return 0

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = MANUAL_JAVA.with_suffix(f".java.{stamp}.bak")
    shutil.copy2(MANUAL_JAVA, backup)
    MANUAL_JAVA.write_text(new_src, encoding="utf-8")
    info(f"已写回 {MANUAL_JAVA}")
    info(f"  备份：{backup}")
    info("  下一步：cd MDK-1.21.1-ModDevGradle-main && gradlew jar   （或跑 HuhuTalk-Demo\\_mmbuild.bat）")
    return 0


def do_dump(docx_path: Path) -> int:
    """把 Word 里读到的结构原样打印出来（审阅用），并单独列出所有"给 ds / 写给 deepseek"指令标记。"""
    doc_sections = read_docx(docx_path)
    total = 0
    markers: list[str] = []
    for sec in doc_sections:
        info(f'=== 大栏目 [{sec["src_id"]}] {sec["title"]}')
        for en in sec["entries"]:
            total += 1
            eid = en["src_id"]
            tag = f'[{eid[0]}-{eid[1]}]' if eid else "[无编号]"
            info(f'  -- {tag} {en["title"]}')
            for line in en["body"].split("\n"):
                info(f'       | {line}')
            if en.get("image"):
                info(f'       [配图 {figure_label(en["image"])}]')
            blob = en["title"] + "\n" + en["body"]
            for m in re.finditer(r"(给\s*ds[^）)]*|写给\s*deepseek[^）)]*|给deepseek[^）)]*)", blob, re.I):
                markers.append(f'{tag} {en["title"]} → {m.group(0).strip()}')
    info()
    info(f"共 {len(doc_sections)} 节 / {total} 条")
    if markers:
        info()
        info("=== 指令标记（要照做，且不该进模组正文）===")
        for mk in markers:
            info("  " + mk)
    return 0


def main() -> int:
    global MANUAL_JAVA
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        sys.stderr.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="HuhuTalk 手册 ⇄ Word 往返工具")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_exp = sub.add_parser("export", help="从 ManualPage.java 导出 Word")
    p_exp.add_argument("-o", "--out", default=str(DEFAULT_DOCX))
    p_exp.add_argument("--manual", default=str(MANUAL_JAVA), help="源文件（默认 ManualPage.java）")

    p_imp = sub.add_parser("import", help="把改过的 Word 写回 ManualPage.java")
    p_imp.add_argument("-i", "--in", dest="inp", default=str(DEFAULT_DOCX))
    p_imp.add_argument("--manual", default=str(MANUAL_JAVA), help="源文件（默认 ManualPage.java）")
    p_imp.add_argument("--dry-run", action="store_true", help="只打印改动清单与 diff，不写盘")
    p_imp.add_argument("--allow-delete", action="store_true", help="允许删掉文档里不存在的条目")

    p_dump = sub.add_parser("dump", help="只把 Word 结构打印出来（含指令标记扫描）")
    p_dump.add_argument("-i", "--in", dest="inp", default=str(DEFAULT_DOCX))

    sub.add_parser("selftest", help="在临时副本上跑一遍导出→改→回写→幂等自检（不动真文件）")

    args = ap.parse_args()
    if args.cmd == "export":
        MANUAL_JAVA = Path(args.manual)
        return do_export(Path(args.out))
    if args.cmd == "import":
        MANUAL_JAVA = Path(args.manual)
        return do_import(Path(args.inp), args.dry_run, args.allow_delete)
    if args.cmd == "dump":
        return do_dump(Path(args.inp))
    return do_selftest()


# ============================================================================
# 自检：整条往返链路（只在临时副本上跑，绝不动真文件）
# ============================================================================


def do_selftest() -> int:
    import tempfile

    from docx import Document

    global MANUAL_JAVA
    real = MANUAL_JAVA
    tmp = Path(tempfile.mkdtemp(prefix="manual-docx-selftest-"))
    copy = tmp / "ManualPage.java"
    shutil.copy2(real, copy)
    docx = tmp / "manual.docx"
    docx_edited = tmp / "manual-edited.docx"

    MANUAL_JAVA = copy
    info(f"自检目录：{tmp}")
    do_export(docx)

    # ---- 模拟"需求方在 Word 里编辑" ----
    NEW_BODY = [
        '自检改动 A：中文标点（“引号”、破折号——）、反斜杠 \\ 与引号 " 都要原样回来',
        "自检第二段（换行 = 模组里的一行）",
    ]
    doc = Document(str(docx))
    paras = doc.paragraphs
    idx = {p.text.strip(): i for i, p in enumerate(paras)}
    # 1) 改正文：[1-2] 这一条的正文整段换成 NEW_BODY
    start = idx["[1-2] 关于加载"]
    i = start + 1
    while i < len(paras) and _heading_level(paras[i]) == 0:
        i += 1
    for p in paras[start + 1:i]:
        p.text = ""
    body_ps = paras[start + 1:i]
    body_ps[0].text = NEW_BODY[0]
    if len(body_ps) > 1:
        body_ps[1].text = NEW_BODY[1]
    else:
        ins = paras[i].insert_paragraph_before(NEW_BODY[1])
        ins.style = doc.styles["Normal"]
    # 2) 改大栏目名
    paras[idx["[3] 大栏目：关于设置"]].text = "[3] 大栏目：关于设置（改名自检）"
    # 3) 新增一条（编号写成 [新]）→ 追加到 [2] 大栏目末尾（= 插在 [3] 的标题之前）
    anchor = idx["[3] 大栏目：关于设置"]
    p = paras[anchor].insert_paragraph_before("[新] 自检新增条目")
    p.style = doc.styles["Heading 2"]
    p2 = paras[anchor].insert_paragraph_before("自检新增正文第一行")
    p2.style = doc.styles["Normal"]
    p3 = paras[anchor].insert_paragraph_before("自检新增正文第二行")
    p3.style = doc.styles["Normal"]
    # 4) 删掉一条（已关机）→ 需要 --allow-delete
    #    用标题查找而不是写死编号：手册一改结构（这次「联系人设置」并进「联系人」就是）编号就变了。
    d0 = next(i for i, p in enumerate(paras) if p.text.strip().endswith("已关机"))
    j = d0 + 1
    while j < len(paras) and _heading_level(paras[j]) == 0:
        j += 1
    for k in range(j - 1, d0 - 1, -1):
        paras[k]._element.getparent().remove(paras[k]._element)
    doc.save(str(docx_edited))

    before = parse_manual(copy.read_text(encoding="utf-8"))
    n_before = sum(len(s["entries"]) for s in before["sections"])

    info("\n== 回写（允许删除） ==")
    rc = do_import(docx_edited, dry_run=False, allow_delete=True)
    if rc != 0:
        return rc

    after = parse_manual(copy.read_text(encoding="utf-8"))
    s1 = after["sections"][0]
    s2 = after["sections"][1]
    s3 = after["sections"][2]
    s6 = next(s for s in after["sections"] if s["title"] == "联系人")
    checks = [
        ("正文改动已落地", s1["entries"][1]["body"] == "\n".join(NEW_BODY)),
        ("大栏目改名已落地", s3["title"] == "关于设置（改名自检）"),
        ("新条目已追加到 [2] 末尾", s2["entries"][-1]["title"] == "自检新增条目"),
        ("新条目正文正确", s2["entries"][-1]["body"] == "自检新增正文第一行\n自检新增正文第二行"),
        ("条目总数不变（删 1 条 + 新增 1 条）", sum(len(s["entries"]) for s in after["sections"]) == n_before),
        ("「联系人」里不再有已关机", all(e["title"] != "已关机" for e in s6["entries"])),
        ("其余条目原样保留", s1["entries"][0]["body"].startswith("狐狐talk是一款")),
    ]
    ok = True
    for name, passed in checks:
        info(f"  [{'OK' if passed else 'FAIL'}] {name}")
        ok = ok and passed

    info("\n== 幂等复查（同一份 Word 再回写一次） ==")
    text_before = copy.read_text(encoding="utf-8")
    do_import(docx_edited, dry_run=False, allow_delete=True)
    text_after = copy.read_text(encoding="utf-8")
    info(f"  [{'OK' if text_before == text_after else 'FAIL'}] 第二次回写不产生任何变化")
    ok = ok and text_before == text_after

    MANUAL_JAVA = real
    info("\n自检结果：" + ("全部通过 ✓" if ok else "有失败项 ✗"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
